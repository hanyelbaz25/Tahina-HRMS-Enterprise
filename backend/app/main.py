import base64
import hashlib
import hmac
import json
import os
import secrets
import time
import uuid
from datetime import date, datetime, time as dt_time, timedelta
from io import BytesIO
from pathlib import Path

from collections.abc import Callable
from decimal import Decimal

from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy import func, inspect, or_, text
from sqlalchemy.orm import Session
from pydantic import BaseModel
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from app.database import Base, engine, get_db
from app.models import AuditLog, Attendance, Branch, Employee, EmployeeCharge, EmployeeDocument, LeaveRequest, Payroll, PayrollPeriod, User

Base.metadata.create_all(bind=engine)
if "employees" in inspect(engine).get_table_names() and "allowances" not in {x["name"] for x in inspect(engine).get_columns("employees")}:
    with engine.begin() as connection:
        connection.execute(text("ALTER TABLE employees ADD COLUMN allowances NUMERIC(12, 2) DEFAULT 0"))
if "payroll" in inspect(engine).get_table_names():
    payroll_columns = {x["name"] for x in inspect(engine).get_columns("payroll")}
    with engine.begin() as connection:
        for name, definition in {"work_days":"NUMERIC(6,2) DEFAULT 30", "addition_days":"NUMERIC(6,2) DEFAULT 0", "addition_amount":"NUMERIC(12,2) DEFAULT 0", "deduction_days":"NUMERIC(6,2) DEFAULT 0", "deduction_amount":"NUMERIC(12,2) DEFAULT 0"}.items():
            if name not in payroll_columns: connection.execute(text(f"ALTER TABLE payroll ADD COLUMN {name} {definition}"))
        if "branch_id" not in payroll_columns:
            connection.execute(text("ALTER TABLE payroll ADD COLUMN branch_id INTEGER"))
            connection.execute(text("UPDATE payroll SET branch_id = (SELECT branch_id FROM employees WHERE employees.id = payroll.employee_id)"))
            if engine.dialect.name == "postgresql":
                connection.execute(text("ALTER TABLE payroll DROP CONSTRAINT IF EXISTS uq_payroll_employee_month"))
                connection.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS uq_payroll_employee_month_branch ON payroll(employee_id, payroll_month, branch_id)"))
if "employee_charges" in inspect(engine).get_table_names() and "branch_id" not in {x["name"] for x in inspect(engine).get_columns("employee_charges")}:
    with engine.begin() as connection:
        connection.execute(text("ALTER TABLE employee_charges ADD COLUMN branch_id INTEGER"))
        connection.execute(text("UPDATE employee_charges SET branch_id = (SELECT branch_id FROM employees WHERE employees.id = employee_charges.employee_id)"))

app = FastAPI(title="Tahina HRMS Enterprise API", version="2.1.0")
security = HTTPBearer(auto_error=False)
TOKEN_SECRET = os.getenv("TOKEN_SECRET", "change-this-secret-before-production")
TOKEN_TTL = 8 * 60 * 60
UPLOAD_DIR = Path(__file__).parent.parent / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

ROLE_LABELS = {
    "system_admin": "مدير النظام",
    "hr_manager": "مسؤول الموارد البشرية",
}


class LoginRequest(BaseModel):
    username: str
    password: str


class UserInput(BaseModel):
    username: str
    full_name: str
    password: str
    role: str
    is_active: bool = True


class UserUpdate(BaseModel):
    full_name: str
    password: str = ""
    role: str
    is_active: bool = True


class BranchInput(BaseModel):
    name_ar: str
    city: str
    is_active: bool = True


class EmployeeInput(BaseModel):
    full_name_ar: str
    identity_no: str
    phone: str = ""
    job_title: str
    basic_salary: Decimal = Decimal("0")
    allowances: Decimal = Decimal("0")
    branch_id: int
    is_active: bool = True


class AttendanceInput(BaseModel):
    employee_id: int
    work_date: date
    check_in: dt_time | None = None
    check_out: dt_time | None = None
    notes: str = ""


class PayrollInput(BaseModel):
    employee_id: int
    branch_id: int
    payroll_month: str
    work_days: Decimal = Decimal("30")
    addition_days: Decimal = Decimal("0")
    addition_amount: Decimal = Decimal("0")
    deduction_days: Decimal = Decimal("0")
    deduction_amount: Decimal = Decimal("0")
    notes: str = ""


class EmployeeChargeInput(BaseModel):
    employee_id: int
    branch_id: int
    charge_date: date
    charge_type: str
    amount: Decimal
    deduction_start_month: str
    installment_count: int = 1
    funding_source: str = "branch"
    notes: str = ""


class LeaveInput(BaseModel):
    employee_id: int
    leave_type: str
    start_date: date
    end_date: date
    reason: str = ""


def hash_password(password: str, salt: str | None = None) -> str:
    salt_bytes = bytes.fromhex(salt) if salt else secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt_bytes, 310_000)
    return f"{salt_bytes.hex()}:{digest.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        salt, expected = stored.split(":", 1)
    except ValueError:
        return False
    actual = hash_password(password, salt).split(":", 1)[1]
    return hmac.compare_digest(actual, expected)


def make_token(user: User) -> str:
    payload = {"sub": user.username, "role": user.role, "exp": int(time.time()) + TOKEN_TTL}
    body = base64.urlsafe_b64encode(json.dumps(payload).encode()).decode().rstrip("=")
    signature = hmac.new(TOKEN_SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
    return f"{body}.{signature}"


def current_user(
    credentials: HTTPAuthorizationCredentials | None = Depends(security),
    db: Session = Depends(get_db),
) -> User:
    if not credentials:
        raise HTTPException(status_code=401, detail="يلزم تسجيل الدخول")
    try:
        body, signature = credentials.credentials.split(".", 1)
        expected = hmac.new(TOKEN_SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(signature, expected):
            raise ValueError
        padded = body + "=" * (-len(body) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded))
        if payload["exp"] < time.time():
            raise ValueError
    except (ValueError, KeyError, json.JSONDecodeError):
        raise HTTPException(status_code=401, detail="جلسة الدخول غير صالحة أو منتهية")
    user = db.query(User).filter(User.username == payload["sub"], User.is_active.is_(True)).first()
    if not user:
        raise HTTPException(status_code=401, detail="المستخدم غير متاح")
    return user


def require_roles(*roles: str) -> Callable:
    def checker(user: User = Depends(current_user)) -> User:
        if user.role not in roles:
            raise HTTPException(status_code=403, detail="ليس لديك صلاحية لتنفيذ هذا الإجراء")
        return user
    return checker


def user_dict(item: User) -> dict:
    return {"id": item.id, "username": item.username, "full_name": item.full_name,
            "role": item.role, "role_label": ROLE_LABELS.get(item.role, item.role),
            "is_active": item.is_active, "created_at": item.created_at.isoformat() if item.created_at else None}


def branch_dict(branch: Branch) -> dict:
    return {"id": branch.id, "code": branch.code, "name_ar": branch.name_ar, "city": branch.city, "is_active": branch.is_active}


def employee_dict(employee: Employee) -> dict:
    return {
        "id": employee.id, "employee_no": employee.employee_no, "full_name_ar": employee.full_name_ar,
        "identity_no": employee.identity_no, "phone": employee.phone, "job_title": employee.job_title,
        "basic_salary": float(employee.basic_salary), "allowances": float(employee.allowances), "branch_id": employee.branch_id,
        "branch_name": employee.branch.name_ar, "is_active": employee.is_active,
    }


def attendance_dict(item: Attendance) -> dict:
    return {"id": item.id, "employee_id": item.employee_id, "employee_no": item.employee.employee_no,
            "employee_name": item.employee.full_name_ar, "work_date": item.work_date.isoformat(),
            "check_in": item.check_in.strftime("%H:%M") if item.check_in else None,
            "check_out": item.check_out.strftime("%H:%M") if item.check_out else None,
            "notes": item.notes, "source": item.source}


def month_bounds(month: str) -> tuple[date, date]:
    start = datetime.strptime(month, "%Y-%m").date().replace(day=1)
    end = date(start.year + (start.month == 12), 1 if start.month == 12 else start.month + 1, 1)
    return start, end


def month_index(month: str) -> int:
    value = datetime.strptime(month, "%Y-%m")
    return value.year * 12 + value.month - 1


def scheduled_charge_amount(item: EmployeeCharge, month: str) -> Decimal:
    offset = month_index(month) - month_index(item.deduction_start_month)
    if offset < 0 or offset >= item.installment_count: return Decimal("0")
    part = (Decimal(item.amount) / Decimal(item.installment_count)).quantize(Decimal("0.01"))
    if offset == item.installment_count - 1:
        return Decimal(item.amount) - part * Decimal(item.installment_count - 1)
    return part


def audit(db: Session, user: User, action: str, entity: str, details: str) -> None:
    db.add(AuditLog(username=user.username, action=action, entity=entity, details=details[:500]))


def overlap_days(start: date, end: date, range_start: date, range_end: date) -> int:
    overlap_start = max(start, range_start)
    overlap_end = min(end, range_end - timedelta(days=1))
    return max((overlap_end - overlap_start).days + 1, 0)


def payroll_dict(item: Payroll, db: Session | None = None) -> dict:
    monthly_basic = Decimal(item.employee.basic_salary); monthly_allowances = Decimal(item.employee.allowances)
    day_value = monthly_basic / Decimal("30")
    basic = day_value * Decimal(item.work_days)
    allowances = (monthly_allowances / Decimal("30")) * Decimal(item.work_days)
    total_additions = Decimal(item.addition_amount) + Decimal(item.addition_days) * day_value
    total_deductions = Decimal(item.deduction_amount) + Decimal(item.deduction_days) * day_value
    charges = Decimal("0")
    rewards = Decimal("0")
    leave_deduction_days = 0
    approved = False
    if db:
        employee_charges = db.query(EmployeeCharge).filter(EmployeeCharge.employee_id == item.employee_id, EmployeeCharge.branch_id == item.branch_id).all()
        charges = sum((scheduled_charge_amount(x, item.payroll_month) for x in employee_charges if x.charge_type != "reward"), Decimal("0"))
        rewards = sum((scheduled_charge_amount(x, item.payroll_month) for x in employee_charges if x.charge_type == "reward"), Decimal("0"))
        period = payroll_period(db, item.payroll_month)
        approved = bool(period and period.status == "approved")
        start, end = month_bounds(item.payroll_month)
        unpaid_leaves = db.query(LeaveRequest).filter(
            LeaveRequest.employee_id == item.employee_id,
            LeaveRequest.status == "approved",
            LeaveRequest.leave_type.in_(("unpaid", "unexcused")),
            LeaveRequest.start_date < end,
            LeaveRequest.end_date >= start,
        ).all()
        leave_deduction_days = sum(overlap_days(x.start_date, x.end_date, start, end) for x in unpaid_leaves) if item.branch_id == item.employee.branch_id else 0
        total_deductions += Decimal(leave_deduction_days) * day_value
    applied_charges = charges if approved else Decimal("0")
    applied_rewards = rewards if approved else Decimal("0")
    net = basic + allowances + total_additions + applied_rewards - total_deductions - applied_charges
    return {"id": item.id, "employee_id": item.employee_id, "employee_no": item.employee.employee_no,
            "employee_name": item.employee.full_name_ar, "branch_id": item.branch_id, "branch_name": item.branch.name_ar,
            "job_title": item.employee.job_title, "payroll_month": item.payroll_month,
            "monthly_basic_salary": float(monthly_basic), "monthly_allowances": float(monthly_allowances),
            "basic_salary": float(basic), "allowances": float(allowances), "work_days": float(item.work_days),
            "addition_days": float(item.addition_days), "addition_amount": float(item.addition_amount), "additions": float(total_additions),
            "deduction_days": float(item.deduction_days), "leave_deduction_days": leave_deduction_days,
            "deduction_amount": float(item.deduction_amount), "deductions": float(total_deductions),
            "day_value": float(day_value), "daily_charges": float(charges), "applied_daily_charges": float(applied_charges),
            "daily_rewards": float(rewards), "applied_daily_rewards": float(applied_rewards),
            "net_salary": float(net), "notes": item.notes, "source": item.source}


def payroll_period(db: Session, month: str) -> PayrollPeriod | None:
    return db.query(PayrollPeriod).filter(PayrollPeriod.payroll_month == month).first()


def ensure_payroll_open(db: Session, month: str) -> None:
    period = payroll_period(db, month)
    if period and period.status == "approved":
        raise HTTPException(status_code=409, detail="مسير هذا الشهر معتمد ومغلق أمام التعديل")


def parse_excel_date(value) -> date:
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def parse_excel_time(value) -> dt_time | None:
    if value in (None, ""): return None
    if isinstance(value, datetime): return value.time().replace(second=0, microsecond=0)
    if isinstance(value, dt_time): return value.replace(second=0, microsecond=0)
    return datetime.strptime(str(value).strip(), "%H:%M").time()


def parse_payroll_month(value) -> str:
    if isinstance(value, datetime): return value.strftime("%Y-%m")
    if isinstance(value, date): return value.strftime("%Y-%m")
    if isinstance(value, (int, float)): return from_excel(value).strftime("%Y-%m")
    text_value = str(value).strip().lstrip("'")
    datetime.strptime(text_value, "%Y-%m")
    return text_value


def seed_users() -> None:
    defaults = [
        (os.getenv("ADMIN_USERNAME", "admin"), os.getenv("ADMIN_PASSWORD", "Admin@123"), "مدير النظام", "system_admin"),
        (os.getenv("HR_USERNAME", "hr"), os.getenv("HR_PASSWORD", "HR@12345"), "مسؤول الموارد البشرية", "hr_manager"),
    ]
    with Session(engine) as db:
        for username, password, full_name, role in defaults:
            if not db.query(User).filter(User.username == username).first():
                db.add(User(username=username, full_name=full_name, password_hash=hash_password(password), role=role))
        db.commit()


seed_users()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
def root():
    return {"message": "Tahina HRMS Enterprise API is running"}

@app.get("/health")
def health():
    with engine.connect() as connection:
        connection.execute(text("SELECT 1"))
    return {"status": "healthy", "database": "connected"}


@app.post("/api/v1/auth/login")
def login(data: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.username == data.username.strip(), User.is_active.is_(True)).first()
    if not user or not verify_password(data.password, user.password_hash):
        raise HTTPException(status_code=401, detail="اسم المستخدم أو كلمة المرور غير صحيحة")
    return {
        "access_token": make_token(user),
        "token_type": "bearer",
        "user": {"username": user.username, "full_name": user.full_name, "role": user.role, "role_label": ROLE_LABELS[user.role]},
    }


@app.get("/api/v1/auth/me")
def me(user: User = Depends(current_user)):
    return {"username": user.username, "full_name": user.full_name, "role": user.role, "role_label": ROLE_LABELS[user.role]}


@app.get("/api/v1/users")
def list_users(db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    return [user_dict(item) for item in db.query(User).order_by(User.id).all()]


@app.post("/api/v1/users", status_code=201)
def create_user(data: UserInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    username = data.username.strip().lower()
    if len(username) < 3: raise HTTPException(status_code=400, detail="اسم المستخدم يجب ألا يقل عن 3 أحرف")
    if len(data.password) < 8: raise HTTPException(status_code=400, detail="كلمة المرور يجب ألا تقل عن 8 أحرف")
    if data.role not in ROLE_LABELS: raise HTTPException(status_code=400, detail="الصلاحية غير صحيحة")
    if db.query(User).filter(User.username == username).first(): raise HTTPException(status_code=409, detail="اسم المستخدم مستخدم بالفعل")
    item = User(username=username, full_name=data.full_name.strip(), password_hash=hash_password(data.password), role=data.role, is_active=data.is_active)
    db.add(item); audit(db, user, "create", "user", username); db.commit(); db.refresh(item)
    return user_dict(item)


@app.put("/api/v1/users/{user_id}")
def update_user(user_id: int, data: UserUpdate, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(User, user_id)
    if not item: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if data.role not in ROLE_LABELS: raise HTTPException(status_code=400, detail="الصلاحية غير صحيحة")
    if data.password and len(data.password) < 8: raise HTTPException(status_code=400, detail="كلمة المرور يجب ألا تقل عن 8 أحرف")
    if item.id == user.id and not data.is_active: raise HTTPException(status_code=409, detail="لا يمكنك إيقاف حسابك الحالي")
    item.full_name, item.role, item.is_active = data.full_name.strip(), data.role, data.is_active
    if data.password: item.password_hash = hash_password(data.password)
    audit(db, user, "update", "user", item.username); db.commit(); db.refresh(item)
    return user_dict(item)


@app.delete("/api/v1/users/{user_id}", status_code=204)
def delete_user(user_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(User, user_id)
    if not item: raise HTTPException(status_code=404, detail="المستخدم غير موجود")
    if item.id == user.id: raise HTTPException(status_code=409, detail="لا يمكنك حذف حسابك الحالي")
    audit(db, user, "delete", "user", item.username); db.delete(item); db.commit()

@app.get("/api/v1/dashboard")
def dashboard(user: User = Depends(current_user), db: Session = Depends(get_db)):
    employees_total = db.query(func.count(Employee.id)).filter(Employee.is_active.is_(True)).scalar() or 0
    present_today = db.query(func.count(Attendance.id)).filter(Attendance.work_date == date.today(), Attendance.check_in.isnot(None)).scalar() or 0
    return {
        "employees_total": employees_total,
        "branches_total": db.query(func.count(Branch.id)).scalar() or 0,
        "employees_present_today": present_today,
        "employees_absent_today": max(employees_total - present_today, 0),
        "contracts_expiring": db.query(func.count(EmployeeDocument.id)).filter(EmployeeDocument.document_type == "contract", EmployeeDocument.expiry_date >= date.today(), EmployeeDocument.expiry_date <= date.today() + timedelta(days=90)).scalar() or 0,
        "iqamas_expiring": db.query(func.count(EmployeeDocument.id)).filter(EmployeeDocument.document_type == "iqama", EmployeeDocument.expiry_date >= date.today(), EmployeeDocument.expiry_date <= date.today() + timedelta(days=90)).scalar() or 0,
        "payroll_status": "غير مُنشأ",
        "system_version": "2.1.0"
    }


@app.get("/api/v1/branches")
def list_branches(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return [branch_dict(item) for item in db.query(Branch).order_by(Branch.name_ar).all()]


@app.post("/api/v1/branches", status_code=201)
def create_branch(data: BranchInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    branch = Branch(code=f"PENDING-{secrets.token_hex(4)}", **data.model_dump())
    db.add(branch); db.flush()
    branch.code = f"BR{branch.id:03d}"
    db.commit(); db.refresh(branch)
    return branch_dict(branch)


@app.put("/api/v1/branches/{branch_id}")
def update_branch(branch_id: int, data: BranchInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    branch = db.get(Branch, branch_id)
    if not branch: raise HTTPException(status_code=404, detail="الفرع غير موجود")
    for key, value in data.model_dump().items(): setattr(branch, key, value)
    db.commit(); db.refresh(branch)
    return branch_dict(branch)


@app.delete("/api/v1/branches/{branch_id}", status_code=204)
def delete_branch(branch_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    branch = db.get(Branch, branch_id)
    if not branch: raise HTTPException(status_code=404, detail="الفرع غير موجود")
    if db.query(Employee).filter(Employee.branch_id == branch_id).first():
        raise HTTPException(status_code=409, detail="لا يمكن حذف فرع مرتبط بموظفين")
    db.delete(branch); db.commit()


@app.get("/api/v1/employees")
def list_employees(search: str = Query("", max_length=100), db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(Employee)
    if search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(or_(Employee.full_name_ar.ilike(term), Employee.employee_no.ilike(term), Employee.identity_no.ilike(term)))
    return [employee_dict(item) for item in query.order_by(Employee.id.desc()).all()]


@app.post("/api/v1/employees", status_code=201)
def create_employee(data: EmployeeInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not db.get(Branch, data.branch_id): raise HTTPException(status_code=400, detail="الفرع المحدد غير موجود")
    if db.query(Employee).filter(Employee.identity_no == data.identity_no.strip()).first():
        raise HTTPException(status_code=409, detail="رقم الهوية مستخدم بالفعل")
    employee = Employee(employee_no=f"PENDING-{secrets.token_hex(4)}", **data.model_dump())
    db.add(employee); db.flush()
    employee.employee_no = f"EMP{employee.id:04d}"
    db.commit(); db.refresh(employee)
    return employee_dict(employee)


@app.put("/api/v1/employees/{employee_id}")
def update_employee(employee_id: int, data: EmployeeInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    employee = db.get(Employee, employee_id)
    if not employee: raise HTTPException(status_code=404, detail="الموظف غير موجود")
    if not db.get(Branch, data.branch_id): raise HTTPException(status_code=400, detail="الفرع المحدد غير موجود")
    duplicate = db.query(Employee).filter(Employee.id != employee_id, Employee.identity_no == data.identity_no.strip()).first()
    if duplicate: raise HTTPException(status_code=409, detail="رقم الهوية مستخدم بالفعل")
    for key, value in data.model_dump().items(): setattr(employee, key, value)
    db.commit(); db.refresh(employee)
    return employee_dict(employee)


@app.delete("/api/v1/employees/{employee_id}", status_code=204)
def delete_employee(employee_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    employee = db.get(Employee, employee_id)
    if not employee: raise HTTPException(status_code=404, detail="الموظف غير موجود")
    db.delete(employee); db.commit()


@app.post("/api/v1/employees/import")
async def import_employees(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="يرجى رفع ملف Excel بصيغة xlsx")
    try:
        workbook = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = workbook.active
    except Exception:
        raise HTTPException(status_code=400, detail="تعذر قراءة ملف Excel")
    expected = ["كود الموظف", "اسم الموظف", "رقم الهوية", "الراتب الأساسي", "البدلات", "الوظيفة", "الفرع"]
    headers = [str(x.value or "").strip() for x in sheet[1][:7]]
    if headers != expected: raise HTTPException(status_code=400, detail="عناوين الأعمدة غير مطابقة: " + "، ".join(expected))
    created = updated = branches_created = 0; errors: list[dict] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row[:7]): continue
        try:
            name = str(row[1] or "").strip(); identity = str(row[2] or "").strip(); job = str(row[5] or "").strip(); branch_name = str(row[6] or "").strip()
            if not all((name, identity, job, branch_name)): raise ValueError("الاسم والهوية والوظيفة والفرع بيانات إلزامية")
            branch = db.query(Branch).filter(func.lower(Branch.name_ar) == branch_name.lower()).first()
            if not branch:
                branch = Branch(code=f"PENDING-{secrets.token_hex(4)}", name_ar=branch_name, city=branch_name, is_active=True); db.add(branch); db.flush(); branch.code=f"BR{branch.id:03d}"; branches_created += 1
            employee = db.query(Employee).filter(Employee.identity_no == identity).first()
            values = {"full_name_ar": name, "identity_no": identity, "job_title": job, "branch_id": branch.id,
                      "basic_salary": Decimal(str(row[3] or 0)), "allowances": Decimal(str(row[4] or 0)), "is_active": True}
            if employee:
                if user.role != "system_admin": raise ValueError("السجل موجود بالفعل ومسؤول الموارد البشرية لا يملك صلاحية تعديله")
                for key, value in values.items(): setattr(employee, key, value)
                updated += 1
            else:
                employee = Employee(employee_no=f"PENDING-{secrets.token_hex(4)}", phone="", **values); db.add(employee); db.flush(); employee.employee_no=f"EMP{employee.id:04d}"; created += 1
            db.commit()
        except Exception as exc:
            db.rollback(); errors.append({"row": row_number, "message": str(exc)})
    return {"created": created, "updated": updated, "branches_created": branches_created, "errors": errors[:100], "errors_total": len(errors)}


@app.get("/api/v1/attendance")
def list_attendance(work_date: date | None = None, db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(Attendance)
    if work_date: query = query.filter(Attendance.work_date == work_date)
    return [attendance_dict(x) for x in query.order_by(Attendance.work_date.desc(), Attendance.id.desc()).limit(500).all()]


@app.post("/api/v1/attendance/manual")
def save_manual_attendance(data: AttendanceInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not db.get(Employee, data.employee_id): raise HTTPException(status_code=400, detail="الموظف غير موجود")
    item = db.query(Attendance).filter(Attendance.employee_id == data.employee_id, Attendance.work_date == data.work_date).first()
    if item:
        if user.role != "system_admin": raise HTTPException(status_code=403, detail="مسؤول الموارد البشرية لا يملك صلاحية تعديل سجل موجود")
        for key, value in data.model_dump().items(): setattr(item, key, value)
        item.source = "manual"
    else:
        item = Attendance(**data.model_dump(), source="manual"); db.add(item)
    db.commit(); db.refresh(item)
    return attendance_dict(item)


@app.delete("/api/v1/attendance/{attendance_id}", status_code=204)
def delete_attendance(attendance_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(Attendance, attendance_id)
    if not item: raise HTTPException(status_code=404, detail="السجل غير موجود")
    db.delete(item); db.commit()


@app.post("/api/v1/attendance/import")
async def import_attendance(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="يرجى رفع ملف Excel بصيغة xlsx")
    try:
        workbook = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True)
        sheet = workbook.active
    except Exception:
        raise HTTPException(status_code=400, detail="تعذر قراءة ملف Excel")
    expected = ["employee_no", "work_date", "check_in", "check_out", "notes"]
    headers = [str(x.value or "").strip().lower() for x in sheet[1][:5]]
    if headers != expected: raise HTTPException(status_code=400, detail="عناوين الأعمدة غير مطابقة للقالب")
    imported = updated = 0; errors: list[dict] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row[:5]): continue
        try:
            employee_no = str(row[0]).strip(); employee = db.query(Employee).filter(Employee.employee_no == employee_no).first()
            if not employee: raise ValueError(f"كود الموظف {employee_no} غير موجود")
            work_date = parse_excel_date(row[1]); check_in = parse_excel_time(row[2]); check_out = parse_excel_time(row[3]); notes = str(row[4] or "").strip()
            item = db.query(Attendance).filter(Attendance.employee_id == employee.id, Attendance.work_date == work_date).first()
            if item:
                if user.role != "system_admin": raise ValueError("السجل موجود بالفعل ومسؤول الموارد البشرية لا يملك صلاحية تعديله")
                item.check_in, item.check_out, item.notes, item.source = check_in, check_out, notes, "excel"; updated += 1
            else:
                db.add(Attendance(employee_id=employee.id, work_date=work_date, check_in=check_in, check_out=check_out, notes=notes, source="excel")); imported += 1
            db.commit()
        except Exception as exc:
            db.rollback(); errors.append({"row": row_number, "message": str(exc)})
    return {"imported": imported, "updated": updated, "errors": errors[:100], "errors_total": len(errors)}


def document_dict(x: EmployeeDocument) -> dict:
    return {"id": x.id, "employee_id": x.employee_id, "employee_no": x.employee.employee_no, "employee_name": x.employee.full_name_ar, "branch_name": x.employee.branch.name_ar, "document_type": x.document_type, "document_no": x.document_no, "start_date": x.start_date.isoformat() if x.start_date else None, "expiry_date": x.expiry_date.isoformat(), "days_remaining": (x.expiry_date - date.today()).days, "notes": x.notes, "has_file": bool(x.stored_file_name), "original_file_name": x.original_file_name}


@app.get("/api/v1/documents")
def list_documents(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return [document_dict(x) for x in db.query(EmployeeDocument).order_by(EmployeeDocument.expiry_date).all()]


@app.post("/api/v1/documents")
async def create_document(employee_id: int = Form(...), document_type: str = Form(...), document_no: str = Form(""), start_date: date | None = Form(None), expiry_date: date = Form(...), notes: str = Form(""), file: UploadFile | None = File(None), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    employee = db.get(Employee, employee_id)
    if not employee: raise HTTPException(status_code=400, detail="الموظف غير موجود")
    if document_type not in ("contract", "iqama"): raise HTTPException(status_code=400, detail="نوع المستند غير صحيح")
    if start_date and expiry_date < start_date: raise HTTPException(status_code=400, detail="تاريخ الانتهاء يجب أن يكون بعد تاريخ البداية")
    stored = original = ""
    if file and file.filename:
        suffix = Path(file.filename).suffix.lower()
        if suffix not in (".pdf", ".png", ".jpg", ".jpeg"): raise HTTPException(status_code=400, detail="المسموح PDF أو PNG أو JPG")
        content = await file.read()
        if len(content) > 10 * 1024 * 1024: raise HTTPException(status_code=400, detail="حجم الملف يجب ألا يتجاوز 10 ميجابايت")
        stored = f"{uuid.uuid4().hex}{suffix}"; original = Path(file.filename).name
        (UPLOAD_DIR / stored).write_bytes(content)
    item = EmployeeDocument(employee_id=employee_id, document_type=document_type, document_no=document_no, start_date=start_date, expiry_date=expiry_date, notes=notes, original_file_name=original, stored_file_name=stored, created_by=user.username)
    db.add(item); audit(db, user, "create", "document", f"{employee.employee_no} - {document_type} - {expiry_date}")
    db.commit(); db.refresh(item); return document_dict(item)


@app.get("/api/v1/documents/{document_id}/file")
def download_document(document_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    item = db.get(EmployeeDocument, document_id)
    if not item or not item.stored_file_name: raise HTTPException(status_code=404, detail="الملف غير موجود")
    path = UPLOAD_DIR / item.stored_file_name
    if not path.exists(): raise HTTPException(status_code=404, detail="الملف غير موجود على الخادم")
    return FileResponse(path, filename=item.original_file_name)


@app.delete("/api/v1/documents/{document_id}", status_code=204)
def delete_document(document_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(EmployeeDocument, document_id)
    if not item: raise HTTPException(status_code=404, detail="المستند غير موجود")
    audit(db, user, "delete", "document", f"id={item.id} - {item.employee.employee_no}")
    if item.stored_file_name:
        path = UPLOAD_DIR / item.stored_file_name
        if path.exists(): path.unlink()
    db.delete(item); db.commit()


@app.get("/api/v1/leaves")
def list_leaves(year: int = Query(..., ge=2000, le=2100), db: Session = Depends(get_db), user: User = Depends(current_user)):
    start, end = date(year, 1, 1), date(year + 1, 1, 1)
    items = db.query(LeaveRequest).filter(LeaveRequest.start_date < end, LeaveRequest.end_date >= start).order_by(LeaveRequest.start_date.desc()).all()
    return [{"id": x.id, "employee_id": x.employee_id, "employee_no": x.employee.employee_no, "employee_name": x.employee.full_name_ar, "branch_name": x.employee.branch.name_ar, "leave_type": x.leave_type, "start_date": x.start_date.isoformat(), "end_date": x.end_date.isoformat(), "return_date": (x.end_date + timedelta(days=1)).isoformat(), "days": x.days, "reason": x.reason, "status": x.status, "approved_by": x.approved_by} for x in items]


@app.get("/api/v1/leaves/balances")
def leave_balances(year: int = Query(..., ge=2000, le=2100), db: Session = Depends(get_db), user: User = Depends(current_user)):
    year_start, year_end = date(year, 1, 1), date(year + 1, 1, 1)
    result = []
    for employee in db.query(Employee).filter(Employee.is_active.is_(True)).order_by(Employee.employee_no).all():
        annual = db.query(LeaveRequest).filter(LeaveRequest.employee_id == employee.id, LeaveRequest.leave_type == "annual", LeaveRequest.status == "approved", LeaveRequest.start_date < year_end, LeaveRequest.end_date >= year_start).all()
        used = sum(overlap_days(x.start_date, x.end_date, year_start, year_end) for x in annual)
        result.append({"employee_id": employee.id, "employee_no": employee.employee_no, "employee_name": employee.full_name_ar, "branch_name": employee.branch.name_ar, "entitlement": 30, "used": used, "remaining": max(30 - used, 0)})
    return result


@app.get("/api/v1/leaves/movement-report")
def leave_movement_report(year: int = Query(..., ge=2000, le=2100), db: Session = Depends(get_db), user: User = Depends(current_user)):
    today = date.today(); year_start, year_end = date(year, 1, 1), date(year + 1, 1, 1)
    items = db.query(LeaveRequest).filter(LeaveRequest.status == "approved", LeaveRequest.start_date < year_end, LeaveRequest.end_date >= year_start).order_by(LeaveRequest.end_date.desc()).all()
    def row(x: LeaveRequest) -> dict:
        return {"employee_no": x.employee.employee_no, "employee_name": x.employee.full_name_ar, "branch_name": x.employee.branch.name_ar, "leave_type": x.leave_type, "start_date": x.start_date.isoformat(), "end_date": x.end_date.isoformat(), "return_date": (x.end_date + timedelta(days=1)).isoformat(), "days": x.days}
    return {"on_leave": [row(x) for x in items if x.start_date <= today <= x.end_date], "returned": [row(x) for x in items if x.end_date < today]}


@app.post("/api/v1/leaves")
def create_leave(data: LeaveInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    employee = db.get(Employee, data.employee_id)
    if not employee: raise HTTPException(status_code=400, detail="الموظف غير موجود")
    if data.leave_type not in ("annual", "sick", "unpaid", "unexcused"): raise HTTPException(status_code=400, detail="نوع الإجازة غير صحيح")
    if data.end_date < data.start_date: raise HTTPException(status_code=400, detail="تاريخ نهاية الإجازة يجب أن يكون بعد تاريخ البداية")
    conflict = db.query(LeaveRequest).filter(LeaveRequest.employee_id == data.employee_id, LeaveRequest.status != "rejected", LeaveRequest.start_date <= data.end_date, LeaveRequest.end_date >= data.start_date).first()
    if conflict: raise HTTPException(status_code=409, detail="توجد إجازة أخرى متداخلة لهذا الموظف")
    days = (data.end_date - data.start_date).days + 1
    item = LeaveRequest(**data.model_dump(), days=days, created_by=user.username)
    db.add(item); audit(db, user, "create", "leave", f"{employee.employee_no} - {data.start_date} إلى {data.end_date}")
    db.commit(); db.refresh(item); return {"id": item.id, "message": "تم تسجيل طلب الإجازة"}


@app.post("/api/v1/leaves/{leave_id}/{action}")
def decide_leave(leave_id: int, action: str, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    item = db.get(LeaveRequest, leave_id)
    if not item: raise HTTPException(status_code=404, detail="طلب الإجازة غير موجود")
    if action not in ("approve", "reject"): raise HTTPException(status_code=400, detail="الإجراء غير صحيح")
    if action == "approve" and item.leave_type == "annual":
        year_start, year_end = date(item.start_date.year, 1, 1), date(item.start_date.year + 1, 1, 1)
        annual = db.query(LeaveRequest).filter(LeaveRequest.employee_id == item.employee_id, LeaveRequest.id != item.id, LeaveRequest.leave_type == "annual", LeaveRequest.status == "approved", LeaveRequest.start_date < year_end, LeaveRequest.end_date >= year_start).all()
        used = sum(overlap_days(x.start_date, x.end_date, year_start, year_end) for x in annual)
        if used + overlap_days(item.start_date, item.end_date, year_start, year_end) > 30: raise HTTPException(status_code=409, detail="رصيد الإجازة السنوية لا يكفي")
    if action == "approve" and item.leave_type in ("unpaid", "unexcused"):
        index = month_index(item.start_date.strftime("%Y-%m")); last = month_index(item.end_date.strftime("%Y-%m"))
        while index <= last:
            ensure_payroll_open(db, f"{index // 12:04d}-{index % 12 + 1:02d}"); index += 1
    item.status = "approved" if action == "approve" else "rejected"; item.approved_by = user.full_name
    audit(db, user, action, "leave", f"id={item.id} - {item.employee.employee_no}")
    db.commit(); return {"message": "تم اعتماد الإجازة" if action == "approve" else "تم رفض الإجازة"}


@app.delete("/api/v1/leaves/{leave_id}", status_code=204)
def delete_leave(leave_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(LeaveRequest, leave_id)
    if not item: raise HTTPException(status_code=404, detail="طلب الإجازة غير موجود")
    if item.status == "approved" and item.leave_type in ("unpaid", "unexcused"):
        ensure_payroll_open(db, item.start_date.strftime("%Y-%m"))
    audit(db, user, "delete", "leave", f"id={item.id}"); db.delete(item); db.commit()


@app.get("/api/v1/employee-charges")
def list_employee_charges(month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    start, end = month_bounds(month)
    all_items = db.query(EmployeeCharge).order_by(EmployeeCharge.charge_date.desc(), EmployeeCharge.id.desc()).all()
    items = [x for x in all_items if start <= x.charge_date < end or scheduled_charge_amount(x, month) > 0]
    return [{
        "id": x.id, "employee_id": x.employee_id, "employee_no": x.employee.employee_no,
        "employee_name": x.employee.full_name_ar, "branch_id": x.branch_id, "branch_name": x.branch.name_ar,
        "charge_date": x.charge_date.isoformat(), "charge_type": x.charge_type,
        "amount": float(x.amount), "deduction_start_month": x.deduction_start_month,
        "installment_count": x.installment_count, "funding_source": x.funding_source,
        "monthly_amount": float(scheduled_charge_amount(x, month)),
        "notes": x.notes, "created_by": x.created_by,
    } for x in items]


@app.post("/api/v1/employee-charges")
def create_employee_charge(data: EmployeeChargeInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    employee = db.get(Employee, data.employee_id)
    if not employee: raise HTTPException(status_code=400, detail="الموظف غير موجود")
    branch = db.get(Branch, data.branch_id)
    if not branch: raise HTTPException(status_code=400, detail="الفرع غير موجود")
    if data.charge_type not in ("advance", "order", "penalty", "reward"): raise HTTPException(status_code=400, detail="نوع الحركة غير صحيح")
    if data.funding_source not in ("branch", "treasury", "bank"): raise HTTPException(status_code=400, detail="مصدر السلفة غير صحيح")
    if data.amount <= 0: raise HTTPException(status_code=400, detail="يجب أن تكون القيمة أكبر من صفر")
    try: datetime.strptime(data.deduction_start_month, "%Y-%m")
    except ValueError: raise HTTPException(status_code=400, detail="صيغة شهر بدء الخصم غير صحيحة")
    if data.installment_count < 1 or data.installment_count > 60: raise HTTPException(status_code=400, detail="عدد الأقساط يجب أن يكون بين 1 و60")
    ensure_payroll_open(db, data.deduction_start_month)
    item = EmployeeCharge(**data.model_dump(), created_by=user.username)
    db.add(item)
    audit(db, user, "create", "employee_charge", f"{employee.employee_no} - {branch.name_ar} - {data.charge_type} - {data.amount}")
    db.commit(); db.refresh(item)
    return {"id": item.id, "message": "تم تسجيل الحركة اليومية"}


@app.delete("/api/v1/employee-charges/{charge_id}", status_code=204)
def delete_employee_charge(charge_id: int, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    item = db.get(EmployeeCharge, charge_id)
    if not item: raise HTTPException(status_code=404, detail="الحركة غير موجودة")
    for offset in range(item.installment_count):
        index = month_index(item.deduction_start_month) + offset
        ensure_payroll_open(db, f"{index // 12:04d}-{index % 12 + 1:02d}")
    audit(db, user, "delete", "employee_charge", f"id={item.id} - {item.amount}")
    db.delete(item); db.commit()


@app.get("/api/v1/audit-logs")
def list_audit_logs(limit: int = Query(100, ge=1, le=500), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    return [{"id": x.id, "username": x.username, "action": x.action, "entity": x.entity, "details": x.details, "created_at": x.created_at.isoformat()} for x in db.query(AuditLog).order_by(AuditLog.id.desc()).limit(limit).all()]


@app.get("/api/v1/employees/{employee_id}/monthly-report")
def employee_monthly_report(employee_id: int, month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    employee = db.get(Employee, employee_id)
    if not employee: raise HTTPException(status_code=404, detail="الموظف غير موجود")
    start, end = month_bounds(month)
    attendance = db.query(Attendance).filter(Attendance.employee_id == employee_id, Attendance.work_date >= start, Attendance.work_date < end).order_by(Attendance.work_date).all()
    charges = db.query(EmployeeCharge).filter(EmployeeCharge.employee_id == employee_id).order_by(EmployeeCharge.charge_date).all()
    relevant_charges = [x for x in charges if scheduled_charge_amount(x, month) > 0 or start <= x.charge_date < end]
    payroll_items = db.query(Payroll).filter(Payroll.employee_id == employee_id, Payroll.payroll_month == month).order_by(Payroll.branch_id).all()
    payrolls = [payroll_dict(x, db) for x in payroll_items]
    payroll = None
    if payrolls:
        payroll = dict(payrolls[0]); payroll["branch_name"] = "جميع الفروع"
        for field in ("basic_salary", "allowances", "work_days", "addition_days", "addition_amount", "additions", "deduction_days", "leave_deduction_days", "deduction_amount", "deductions", "daily_charges", "applied_daily_charges", "daily_rewards", "applied_daily_rewards", "net_salary"):
            payroll[field] = sum(float(x.get(field, 0)) for x in payrolls)
    leaves = db.query(LeaveRequest).filter(LeaveRequest.employee_id == employee_id, LeaveRequest.status == "approved", LeaveRequest.start_date < end, LeaveRequest.end_date >= start).order_by(LeaveRequest.start_date).all()
    return {
        "employee": employee_dict(employee), "month": month,
        "attendance": [attendance_dict(x) for x in attendance],
        "charges": [{"charge_date": x.charge_date.isoformat(), "charge_type": x.charge_type, "branch_name": x.branch.name_ar, "amount": float(x.amount), "funding_source": x.funding_source, "deduction_start_month": x.deduction_start_month, "installment_count": x.installment_count, "month_amount": float(scheduled_charge_amount(x, month)), "notes": x.notes} for x in relevant_charges],
        "leaves": [{"leave_type": x.leave_type, "start_date": x.start_date.isoformat(), "end_date": x.end_date.isoformat(), "days": overlap_days(x.start_date, x.end_date, start, end), "reason": x.reason} for x in leaves],
        "payroll": payroll, "payrolls": payrolls,
    }


@app.get("/api/v1/payroll")
def list_payroll(payroll_month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    return [payroll_dict(x, db) for x in db.query(Payroll).filter(Payroll.payroll_month == payroll_month).order_by(Payroll.id).all()]


@app.post("/api/v1/payroll/manual")
def save_payroll(data: PayrollInput, db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not db.get(Employee, data.employee_id): raise HTTPException(status_code=400, detail="الموظف غير موجود")
    if not db.get(Branch, data.branch_id): raise HTTPException(status_code=400, detail="الفرع غير موجود")
    try: datetime.strptime(data.payroll_month, "%Y-%m")
    except ValueError: raise HTTPException(status_code=400, detail="صيغة الشهر يجب أن تكون YYYY-MM")
    ensure_payroll_open(db, data.payroll_month)
    item = db.query(Payroll).filter(Payroll.employee_id == data.employee_id, Payroll.payroll_month == data.payroll_month, Payroll.branch_id == data.branch_id).first()
    other_days = db.query(func.coalesce(func.sum(Payroll.work_days), 0)).filter(Payroll.employee_id == data.employee_id, Payroll.payroll_month == data.payroll_month, Payroll.id != (item.id if item else 0)).scalar() or 0
    if data.work_days < 0 or Decimal(other_days) + data.work_days > 30: raise HTTPException(status_code=409, detail="إجمالي أيام الموظف في جميع الفروع لا يجوز أن يتجاوز 30 يومًا")
    action = "update" if item else "create"
    if item:
        if user.role != "system_admin": raise HTTPException(status_code=403, detail="مسؤول الموارد البشرية لا يملك صلاحية تعديل مسير موجود")
        for key,value in data.model_dump().items(): setattr(item,key,value)
        item.source = "manual"
    else:
        item = Payroll(**data.model_dump(), source="manual"); db.add(item)
    audit(db, user, action, "payroll", f"{data.payroll_month} - employee_id={data.employee_id}")
    db.commit(); db.refresh(item); return payroll_dict(item, db)


@app.post("/api/v1/payroll/import")
async def import_payroll(file: UploadFile = File(...), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not file.filename or not file.filename.lower().endswith(".xlsx"): raise HTTPException(status_code=400, detail="يرجى رفع ملف Excel بصيغة xlsx")
    try: sheet = load_workbook(BytesIO(await file.read()), read_only=True, data_only=True).active
    except Exception: raise HTTPException(status_code=400, detail="تعذر قراءة ملف Excel")
    expected=["كود الموظف","اسم الموظف","الفرع","شهر المسير","عدد أيام العمل","أيام الإضافي","مبلغ الإضافي","أيام الغياب","مبلغ الخصم","ملاحظات"]
    if [str(x.value or "").strip() for x in sheet[1][:10]] != expected: raise HTTPException(status_code=400, detail="عناوين الأعمدة غير مطابقة لقالب الرواتب")
    imported=updated=0; errors=[]
    for number,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
        if not any(row[:5]): continue
        try:
            employee=db.query(Employee).filter(Employee.employee_no==str(row[0]).strip()).first()
            if not employee: raise ValueError("كود الموظف غير موجود")
            branch=db.query(Branch).filter(Branch.name_ar==str(row[2]).strip()).first()
            if not branch: raise ValueError("اسم الفرع غير موجود")
            month=parse_payroll_month(row[3])
            ensure_payroll_open(db, month)
            item=db.query(Payroll).filter(Payroll.employee_id==employee.id,Payroll.payroll_month==month,Payroll.branch_id==branch.id).first()
            values={"work_days":Decimal(str(row[4] or 0)),"addition_days":Decimal(str(row[5] or 0)),"addition_amount":Decimal(str(row[6] or 0)),"deduction_days":Decimal(str(row[7] or 0)),"deduction_amount":Decimal(str(row[8] or 0)),"notes":str(row[9] or ""),"source":"excel"}
            other_days=db.query(func.coalesce(func.sum(Payroll.work_days),0)).filter(Payroll.employee_id==employee.id,Payroll.payroll_month==month,Payroll.id!=(item.id if item else 0)).scalar() or 0
            if Decimal(other_days)+values["work_days"]>30: raise ValueError("إجمالي أيام الموظف في جميع الفروع يتجاوز 30 يومًا")
            if item:
                if user.role != "system_admin": raise ValueError("المسير موجود بالفعل ومسؤول الموارد البشرية لا يملك صلاحية تعديله")
                for key,value in values.items(): setattr(item,key,value)
                updated+=1
            else: db.add(Payroll(employee_id=employee.id,branch_id=branch.id,payroll_month=month,**values)); imported+=1
            audit(db, user, "excel_import", "payroll", f"{month} - {employee.employee_no}")
            db.commit()
        except Exception as exc: db.rollback(); errors.append({"row":number,"message":str(exc)})
    return {"imported":imported,"updated":updated,"errors_total":len(errors),"errors":errors[:100]}


@app.get("/api/v1/payroll/status")
def get_payroll_status(payroll_month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    period = payroll_period(db, payroll_month)
    return {
        "payroll_month": payroll_month,
        "status": period.status if period else "open",
        "approved_by": period.approved_by if period else "",
        "approved_at": period.approved_at.isoformat() if period and period.approved_at else None,
    }


@app.post("/api/v1/payroll/approve")
def approve_payroll(payroll_month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin", "hr_manager"))):
    if not db.query(Payroll).filter(Payroll.payroll_month == payroll_month).first():
        raise HTTPException(status_code=400, detail="لا يوجد مسير رواتب لاعتماده في هذا الشهر")
    period = payroll_period(db, payroll_month)
    if not period:
        period = PayrollPeriod(payroll_month=payroll_month)
        db.add(period)
    period.status = "approved"
    period.approved_by = user.full_name
    period.approved_at = datetime.now()
    audit(db, user, "approve", "payroll_period", payroll_month)
    db.commit()
    return {"message": "تم اعتماد وإقفال مسير الشهر"}


@app.post("/api/v1/payroll/reopen")
def reopen_payroll(payroll_month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(require_roles("system_admin"))):
    period = payroll_period(db, payroll_month)
    if not period:
        period = PayrollPeriod(payroll_month=payroll_month)
        db.add(period)
    period.status = "open"
    period.approved_by = ""
    period.approved_at = None
    audit(db, user, "reopen", "payroll_period", payroll_month)
    db.commit()
    return {"message": "تمت إعادة فتح مسير الشهر للتعديل"}


@app.get("/api/v1/payroll/export")
def export_payroll(payroll_month: str = Query(..., pattern=r"^\d{4}-\d{2}$"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    items = db.query(Payroll).filter(Payroll.payroll_month == payroll_month).order_by(Payroll.id).all()
    workbook = Workbook(); sheet = workbook.active; sheet.title = "مسير الرواتب"; sheet.sheet_view.rightToLeft = True
    headers = ["كود الموظف","اسم الموظف","الفرع","الوظيفة","الشهر","الراتب الأساسي","البدلات","أيام العمل","أيام الغياب","خصم الغياب","الخصومات المالية","أيام الإضافي","مبلغ الإضافي","إجمالي الإضافي","إجمالي الخصم","السلف والأوردرات والجزاءات","المخصوم بعد الاعتماد","المكافآت المستحقة","المضاف بعد الاعتماد","صافي الراتب","ملاحظات"]
    sheet.merge_cells("B1:U3"); sheet["B1"]="شركة طحينة - مسير الرواتب"; sheet["B1"].font=Font(size=18,bold=True); sheet["B1"].alignment=Alignment(horizontal="center",vertical="center")
    logo_path=Path(__file__).parent / "static" / "tahina-logo.png"
    if logo_path.exists():
        logo=ExcelImage(str(logo_path)); logo.width=105; logo.height=75; sheet.add_image(logo,"A1")
    sheet.append([]); sheet.append(headers)
    for item in items:
        data = payroll_dict(item, db); absence_deduction = data["deduction_days"] * data["day_value"]
        sheet.append([data["employee_no"],data["employee_name"],data["branch_name"],data["job_title"],data["payroll_month"],data["basic_salary"],data["allowances"],data["work_days"],data["deduction_days"],absence_deduction,data["deduction_amount"],data["addition_days"],data["addition_amount"],data["additions"],data["deductions"],data["daily_charges"],data["applied_daily_charges"],data["daily_rewards"],data["applied_daily_rewards"],data["net_salary"],data["notes"]])
    for cell in sheet[5]: cell.fill=PatternFill("solid",fgColor="FF6B00"); cell.font=Font(color="FFFFFF",bold=True); cell.alignment=Alignment(horizontal="center")
    sheet.freeze_panes="A6"; sheet.auto_filter.ref=f"A5:U{sheet.max_row}"
    widths=[15,25,18,20,12,16,14,12,12,15,17,13,15,16,15,20,20,18,18,16,28]
    for index,width in enumerate(widths,1): sheet.column_dimensions[get_column_letter(index)].width=width
    output=BytesIO(); workbook.save(output); output.seek(0)
    headers_out={"Content-Disposition":f'attachment; filename="payroll-{payroll_month}.xlsx"'}
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers=headers_out)
